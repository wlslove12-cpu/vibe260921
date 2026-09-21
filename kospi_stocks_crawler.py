"""네이버 증권 국내 종목 리스트 크롤러 (코스피 거래대금 상위 200 등, 페이징 수집)

https://stock.naver.com/market/stock/kr 의 종목 표(종목명/현재가/전일대비/거래량/거래대금/고가/저가/시가총액)를 읽는다.
이 페이지는 자바스크립트(Next.js)로 표를 그리기 때문에 requests로 받은 HTML에는 표가 없다. 그래서 두 가지 방식을 제공한다.

  1) 기본        : 표를 채우는 JSON API를 페이지 단위로 호출한다(페이징).
                   startIdx 는 '0부터 시작하는 페이지 번호', pageSize 는 '페이지 크기'다.
                   예) pageSize=50 이면 startIdx 0,1,2,3 을 차례로 호출해 200개를 모은다.
  2) --render / --html : 브라우저로 렌더링한 화면(또는 저장해 둔 HTML)의 <table>을 BeautifulSoup으로 파싱한다.
                   화면에는 처음 100행만 있고("항목 더보기"를 눌러야 더 나옴) 시장 탭도 '전체'가 기본이라 200개 수집에는 API 방식을 쓴다.

사용법:
    python kospi_stocks_crawler.py                          # 코스피 거래대금 상위 200개
    python kospi_stocks_crawler.py -n 100 --market KOSDAQ   # 코스닥 상위 100개
    python kospi_stocks_crawler.py --order marketSum        # 시가총액 순 200개 (--help 로 정렬 기준 목록 확인)
    python kospi_stocks_crawler.py --render                 # 렌더링한 화면의 표를 BeautifulSoup으로 파싱(100행)
    python kospi_stocks_crawler.py --html page.html         # 저장해 둔 HTML 파일의 표를 파싱

저장: {이름}.json, {이름}.csv(엑셀에서 한글 정상), {이름}.xlsx(openpyxl 설치 시)

주의: 개인 학습용의 소량 수집을 전제로 했다. 페이지 요청 사이의 지연(DELAY)을 줄이지 말고,
      수집한 데이터를 재배포하기 전에는 네이버 증권의 이용 약관을 확인한다.
"""
import argparse
import csv
import json
import re
import subprocess
import sys
import time

import requests
from bs4 import BeautifulSoup

from kospi200_crawler import find_browser, render_page

PAGE_URL = "https://stock.naver.com/market/stock/kr/stocklist/{order}"
STOCK_API = "https://stock.naver.com/api/domestic/market/stock/default"
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
    "Referer": "https://stock.naver.com/market/stock/kr",
}
TIMEOUT = 10        # 초
DELAY = 0.3         # 페이지 사이의 대기 시간(초)
PAGE_SIZE = 50      # 화면이 쓰는 페이지 크기
MAX_PAGES = 50      # 무한 루프 방지

MARKETS = ("KOSPI", "KOSDAQ", "ALL")
MARKET_NAMES = {"0": "코스피", "1": "코스닥"}   # API의 sosok 값
# 정렬 기준: API의 orderType -> (화면 주소 끝부분, 화면 칩 이름). 둘의 이름이 달라서 대응표를 둔다.
# ('거래량 급증/급감' 칩은 API 값을 찾지 못해 제외했다.)
ORDERS = {
    "priceTop": ("priceTop", "거래대금 상위"),
    "marketSum": ("capitalization", "시가총액"),
    "searchTop": ("top", "인기 종목"),
    "up": ("upper", "상승"),
    "flat": ("flat", "보합"),
    "down": ("lower", "하락"),
    "quantTop": ("trading", "거래량 상위"),
    "high52week": ("high52week", "52주 최고"),
    "low52week": ("low52week", "52주 최저"),
}

FIELDS = ["rank", "code", "name", "market", "price", "change", "change_rate", "volume",
          "trade_value", "open", "high", "low", "market_cap", "status"]
FIELD_LABELS = {
    "rank": "순위", "code": "종목코드", "name": "종목명", "market": "시장", "price": "현재가",
    "change": "전일대비", "change_rate": "등락률(%)", "volume": "거래량(주)", "trade_value": "거래대금(원)",
    "open": "시가", "high": "고가", "low": "저가", "market_cap": "시가총액(원)", "status": "장 상태",
}
DOWN_GB = ("4", "5")   # upDownGb: 1 상한 / 2 상승 / 3 보합 / 4 하한 / 5 하락


def to_number(text):
    """'1,112.16' -> 1112.16, '273000' -> 273000, '5587437000000.0' -> 5587437000000. 숫자가 아니면 None."""
    try:
        value = float(str(text).replace(",", "").replace("+", "").strip())
    except ValueError:
        return None
    return int(value) if value.is_integer() else value


# ---------- 1. JSON API 방식 (기본, 페이징) ----------

def parse_item(item, rank):
    """API 응답 항목 하나를 정리된 dict로 바꾼다. 전일대비는 방향에 맞는 부호로 통일한다."""
    change = to_number(item.get("prevChangePrice"))
    rate = to_number(item.get("prevChangeRate"))
    if item.get("upDownGb") in DOWN_GB:
        change = -abs(change) if change is not None else None
        rate = -abs(rate) if rate is not None else None
    return {
        "rank": rank,
        "code": item["itemcode"],
        "name": item["itemname"],
        "market": MARKET_NAMES.get(item.get("sosok")),
        "price": to_number(item.get("nowPrice")),
        "change": change,
        "change_rate": rate,
        "volume": to_number(item.get("tradeVolume")),
        "trade_value": to_number(item.get("tradeAmount")),      # 원
        "open": to_number(item.get("openPrice")),
        "high": to_number(item.get("highPrice")),
        "low": to_number(item.get("lowPrice")),
        "market_cap": to_number(item.get("marketSum")),          # 원
        "status": item.get("marketStatus"),
    }


def fetch_stocks(session, market="KOSPI", order="priceTop", limit=200, page_size=PAGE_SIZE, log=print):
    """페이지를 0번부터 차례로 넘기며 limit개가 모일 때까지 수집한다.

    순위는 실시간으로 바뀌므로 페이지를 넘기는 사이에 같은 종목이 두 번 나올 수 있다.
    종목코드로 중복을 제거하고, 모자라면 다음 페이지를 더 읽는다.
    """
    stocks, seen = [], set()
    for page in range(MAX_PAGES):
        resp = session.get(STOCK_API, timeout=TIMEOUT, params={
            "tradeType": "KRX", "marketType": market, "orderType": order,
            "startIdx": page, "pageSize": page_size,
        })
        resp.raise_for_status()
        items = resp.json()
        if not items:
            break
        for item in items:
            if item["itemcode"] in seen:
                continue
            seen.add(item["itemcode"])
            stocks.append(parse_item(item, len(stocks) + 1))
        log(f"  {page + 1}페이지: {len(items)}개 받음 (누적 {len(stocks)}개)")
        if len(stocks) >= limit or len(items) < page_size:
            break
        time.sleep(DELAY)
    return stocks[:limit]


# ---------- 2. 렌더링한 HTML + BeautifulSoup 방식 ----------
# 클래스 이름 뒤의 해시(__HI_cb 등)는 배포마다 바뀌므로, 앞부분만 [class*=...] 로 찾는다.

HEADER_KEYS = {
    "종목명": "name", "현재가": "price", "전일대비": "change", "거래량": "volume", "거래대금": "trade_value",
    "고가": "high", "저가": "low", "시가": "open", "시가총액": "market_cap",
}
UNIT_SCALE = {"주": 1, "백만": 1_000_000}          # 셀 안의 단위(<strong>) -> 원/주 환산 배수
CAP_UNITS = (("조", 10 ** 12), ("억", 10 ** 8), ("만", 10 ** 4))


def parse_number_cell(td):
    """<span class="SingleLinePrice_price__">273,000</span><strong>백만</strong> 형태의 셀.
    현재가는 바뀐 자릿수만 <span highlight>로 감싸져 있어 텍스트를 이어 붙여야 한다."""
    price = td.select_one('[class*="SingleLinePrice_price__"]')
    if price is None:
        return to_number(td.get_text("", strip=True))
    value = to_number(price.get_text("", strip=True))
    unit = td.select_one('[class*="SingleLinePrice_unit"]')
    scale = UNIT_SCALE.get(unit.get_text(strip=True), 1) if unit else 1
    return value * scale if value is not None else None


def parse_change_cell(td):
    """전일대비 셀: 금액(▲/▼ 옆 숫자)과 (+4.60%). 방향은 화면에 숨겨진 글자(상승/하락)로 안다."""
    amount = td.select_one('[class*="ModulePriceChange_amount"]')
    percent = td.select_one('[class*="ModulePercent"]')
    if amount is None:
        return None, None
    hint = amount.select_one(".a11y")
    direction = hint.get_text(strip=True) if hint else ""
    if hint:
        hint.extract()   # 금액 텍스트에 '상승'이 섞이지 않게 제거
    change = to_number(amount.get_text("", strip=True))
    m = re.search(r"[+-]?[\d.]+", percent.get_text(strip=True)) if percent else None
    rate = to_number(m.group()) if m else None
    if direction == "하락":
        change = -abs(change) if change is not None else None
        rate = -abs(rate) if rate is not None else None
    return change, rate


def parse_market_cap(text):
    """'1,596조 341억' -> 1596034100000000 (원). '6,374억' -> 637400000000."""
    total, found = 0, False
    for unit, scale in CAP_UNITS:
        m = re.search(rf"([\d,]+)\s*{unit}", text)
        if m:
            total += int(m.group(1).replace(",", "")) * scale
            found = True
    return total if found else to_number(text)


def parse_name_cell(td):
    """종목명 셀: 순위(<span class="index">), 종목명, 로고 주소(Stock005930.svg)에서 종목코드."""
    index = td.select_one("span.index")
    name = td.select_one('[class*="SingleLineText_text__"]')
    logo = td.select_one("img[src]")
    m = re.search(r"Stock(\w+)\.svg", logo["src"]) if logo else None
    return (
        to_number(index.get_text(strip=True)) if index else None,
        name.get_text(strip=True) if name else td.get_text(" ", strip=True),
        m.group(1) if m else None,
    )


def find_stock_table(soup):
    for table in soup.select("table"):
        if any(th.get_text(strip=True) == "종목명" for th in table.select("thead th")):
            return table
    raise ValueError("종목 표(<table>)를 찾지 못했습니다. (페이지 구조가 바뀌었거나 아직 로딩 중일 수 있음)")


def parse_table(html):
    """렌더링된 HTML의 종목 표를 읽는다. 열은 머리글(<th>) 이름으로 찾으므로 '표 설정'으로 열이 바뀌어도 동작한다."""
    table = find_stock_table(BeautifulSoup(html, "html.parser"))
    headers = [th.get_text(strip=True) for th in table.select("thead th")]
    stocks = []
    for n, tr in enumerate(table.select("tbody tr"), start=1):
        cells = tr.find_all("td", recursive=False)
        if len(cells) != len(headers):
            continue
        stock = {key: None for key in FIELDS}
        for header, td in zip(headers, cells):
            key = HEADER_KEYS.get(header)
            if key == "name":
                rank, stock["name"], stock["code"] = parse_name_cell(td)
                stock["rank"] = rank or n
            elif key == "change":
                stock["change"], stock["change_rate"] = parse_change_cell(td)
            elif key == "market_cap":
                stock["market_cap"] = parse_market_cap(td.get_text(" ", strip=True))
            elif key:
                stock[key] = parse_number_cell(td)
        stocks.append(stock)
    if not stocks:
        raise ValueError("표는 찾았지만 행이 없습니다. (아직 로딩 중일 수 있음)")
    return stocks


# ---------- 출력 / 저장 ----------

def fmt_won(value):
    """원 단위 금액을 '1,596조 341억' 처럼 읽기 쉽게."""
    if value is None:
        return "-"
    jo, rest = divmod(int(value), 10 ** 12)
    eok = rest // 10 ** 8
    return f"{jo:,}조 {eok:,}억" if jo else f"{eok:,}억"


def print_stocks(stocks, show=10):
    print(f"\n{'순위':>4} {'종목명':<14}{'현재가':>11}{'전일대비':>10}{'등락률':>10}{'거래대금':>14}{'시가총액':>16}")
    for s in stocks[:show]:
        change = "-" if s["change"] is None else f"{s['change']:+,}"
        rate = "-" if s["change_rate"] is None else f"{s['change_rate']:+.2f}%"
        price = "-" if s["price"] is None else f"{s['price']:,}"
        print(f"{s['rank']:>4} {s['name']:<14}{price:>11}{change:>10}{rate:>10}"
              f"{fmt_won(s['trade_value']):>14}{fmt_won(s['market_cap']):>16}")
    if len(stocks) > show:
        print(f"  … 이하 {len(stocks) - show}개 생략 (총 {len(stocks)}개, 전체는 저장 파일 참고)")


def save_xlsx(path, stocks):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "종목"
    ws.append([FIELD_LABELS[k] for k in FIELDS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
    for s in stocks:
        ws.append([s[k] for k in FIELDS])
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, (int, float)) and FIELDS[cell.column - 1] not in ("rank",):
                cell.number_format = "0.00" if FIELDS[cell.column - 1] == "change_rate" else "#,##0"
    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions
    for col, width in zip("ABCDEFGHIJKLMN", (6, 10, 22, 8, 12, 11, 10, 15, 20, 12, 12, 12, 22, 10)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def save(prefix, stocks):
    with open(f"{prefix}.json", "w", encoding="utf-8") as f:
        json.dump(stocks, f, ensure_ascii=False, indent=2)
    with open(f"{prefix}.csv", "w", encoding="utf-8-sig", newline="") as f:   # utf-8-sig: 엑셀에서 한글 정상
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(stocks)
    saved = [f"{prefix}.json", f"{prefix}.csv"]
    try:
        save_xlsx(f"{prefix}.xlsx", stocks)
        saved.append(f"{prefix}.xlsx")
    except ImportError:
        print("(openpyxl이 없어 .xlsx 저장은 건너뜁니다: pip install openpyxl)")
    print(f"\n저장 완료: {', '.join(saved)}")


def main():
    parser = argparse.ArgumentParser(description="네이버 증권 국내 종목 리스트 크롤러 (페이징)")
    parser.add_argument("-n", "--limit", type=int, default=200, help="수집할 종목 수 (기본: 200)")
    parser.add_argument("--market", choices=MARKETS, default="KOSPI", help="시장 (기본: KOSPI)")
    parser.add_argument("--order", choices=list(ORDERS), default="priceTop",
                        help="정렬 기준 (기본: priceTop). " + ", ".join(f"{k}={v[1]}" for k, v in ORDERS.items()))
    parser.add_argument("--page-size", type=int, default=PAGE_SIZE, help=f"한 번에 받을 개수 (기본: {PAGE_SIZE})")
    parser.add_argument("-o", "--output", default="kospi_stocks", help="저장 파일 이름(확장자 제외)")
    parser.add_argument("--render", action="store_true", help="브라우저로 렌더링한 화면의 표를 BeautifulSoup으로 파싱")
    parser.add_argument("--browser", help="Edge/Chrome 실행 파일 경로 (--render 용, 기본은 자동 탐색)")
    parser.add_argument("--html", help="저장해 둔 HTML 파일의 표를 파싱 (네트워크 사용 안 함)")
    args = parser.parse_args()
    if args.limit < 1 or args.page_size < 1:
        parser.error("--limit 와 --page-size 는 1 이상이어야 합니다.")

    try:
        if args.html:
            with open(args.html, encoding="utf-8") as f:
                stocks = parse_table(f.read())
        elif args.render:
            browser = args.browser or find_browser()
            if not browser:
                raise RuntimeError("Edge/Chrome을 찾지 못했습니다. --browser 로 실행 파일 경로를 알려 주세요.")
            print("브라우저로 페이지를 렌더링하는 중… (10~20초 걸립니다)")
            stocks = parse_table(render_page(PAGE_URL.format(order=ORDERS[args.order][0]), browser))
            print("※ 화면은 '전체' 시장 탭의 처음 100행만 보여 줍니다. 시장 선택과 200개 수집은 기본(API) 방식을 쓰세요.")
        else:
            print(f"{args.market} / {args.order} 상위 {args.limit}개를 {args.page_size}개씩 수집합니다.")
            session = requests.Session()
            session.headers.update(HEADERS)
            stocks = fetch_stocks(session, args.market, args.order, args.limit, args.page_size)
    except (requests.RequestException, RuntimeError, ValueError, OSError, subprocess.TimeoutExpired) as e:
        print(f"수집 실패: {type(e).__name__}: {e}", file=sys.stderr)
        sys.exit(1)

    if not args.html and not args.render and len(stocks) < args.limit:
        print(f"※ 요청한 {args.limit}개보다 적은 {len(stocks)}개만 존재/수집되었습니다.")
    print_stocks(stocks)
    save(args.output, stocks)


if __name__ == "__main__":
    main()
