"""네이버 검색 결과(뉴스 목록)를 파싱하고, 각 기사의 본문까지 크롤링한다.

사용법:
    python news_crawler.py                 # 기본 URL('반도체' 검색), 기사 10개(목록 + 본문)
    python news_crawler.py -n 12           # 기사 12개
    python news_crawler.py --list-only     # 본문은 가져오지 않고 검색 목록만 수집 (요청 1번)
    python news_crawler.py -u "<검색 URL>" -o result

동작 순서:
    1. 검색 결과의 뉴스 목록(.fds-news-item-list-desk)에서 기사 항목을 읽는다.
         - 대표 기사(data-nlog-area="...h.tit")와 그 아래 묶인 관련 기사("...i.tit", "...c.tit")
         - 항목마다: 순위, 제목, 요약(대표 기사만), 언론사, 게시 시각(예: 7시간 전), 원문 URL, 네이버뉴스 URL
    2. 네이버뉴스 URL이 있는 항목은 기사 페이지에서 정확한 작성 시각과 본문을 가져온다.
    3. 결과를 JSON, CSV, 엑셀(.xlsx)로 저장한다. (엑셀은 openpyxl 필요)

목록 태그 구조 (검색 결과 HTML):
    .fds-news-item-list-desk
      └ 대표 기사 묶음
          ├ Profile(data-sds-comp="Profile"): 언론사(.sds-comps-profile-info-title-text),
          │                                    시각(.sds-comps-profile-info-subtext), 네이버뉴스 링크(a[...nav])
          ├ a[data-nlog-area$=".h.tit"]  제목      (href = 원문 URL)
          ├ a[data-nlog-area$=".h.body"] 요약
          └ 관련 기사들: a[data-nlog-area$=".i.tit"] / ".c.tit" + 각자의 Profile
    같은 기사의 태그들은 data-nlog-params 안의 gdid 값이 같아서, 이 값으로 항목을 묶는다.

주의:
    - 개인 학습용의 소량 수집을 전제로 했다. 기사 사이의 지연(DELAY)을 줄이지 말 것.
    - 기사 저작권은 각 언론사에 있으므로 수집한 본문을 재배포하지 않는다.
    - 클래스명 중 fender-ui_..., 영문 난수 클래스는 자주 바뀌므로 쓰지 않는다.
      data-nlog-area / data-sds-comp / sds-comps-* 처럼 의미가 있는 이름만 사용한다.
"""
import argparse
import csv
import json
import re
import time

import requests
from bs4 import BeautifulSoup

SEARCH_URL = (
    "https://search.naver.com/search.naver?where=nexearch&sm=top_hty&fbm=0&ie=utf8"
    "&query=%EB%B0%98%EB%8F%84%EC%B2%B4&ackey=oz6tmm7w"
)
HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
        "(KHTML, like Gecko) Chrome/124.0 Safari/537.36"
    ),
    "Accept-Language": "ko-KR,ko;q=0.9",
}
ARTICLE_URL = re.compile(r"https?://n\.news\.naver\.com/mnews/article/\d+/\d+")
TIMEOUT = 10  # 초
DELAY = 1.0   # 기사 본문 요청 사이의 대기 시간(초)

FIELDS = ["rank", "group", "type", "press", "title", "snippet", "list_time",
          "published", "original_url", "naver_url", "content"]

# 기사 본문(#dic_area) 안에서 기사 내용이 아닌 요소: 스크립트, 사진과 사진 설명, 요약 박스
NOISE_SELECTORS = "script, style, .nbd_a, .nbd_im_w, .end_photo_org, .img_desc, .media_end_summary"


def get_soup(session, url):
    resp = session.get(url, timeout=TIMEOUT)
    resp.raise_for_status()
    resp.encoding = "utf-8"
    return BeautifulSoup(resp.text, "html.parser")


def clean(text):
    """링크마다 숨겨진 '새 창 열림' 문구를 지우고 공백을 정리한다."""
    return " ".join(text.replace("새 창 열림", "").split())


def text_of(element):
    return clean(element.get_text()) if element else ""


# ---------- 1. 검색 결과 목록 파싱 ----------

def parse_list(soup):
    """뉴스 목록에서 대표 기사와 관련 기사를 순위 순서로 반환한다."""
    root = soup.select_one(".fds-news-item-list-desk") or soup
    items, seen, group = [], set(), 0

    for title_a in root.select('a[data-nlog-area$=".tit"]'):
        try:
            params = json.loads(title_a.get("data-nlog-params") or "{}")
        except ValueError:
            params = {}
        gdid = params.get("gdid")
        if not gdid or gdid in seen:
            continue
        seen.add(gdid)

        # 같은 gdid를 가진 요약/네이버뉴스 링크를 찾는다
        nav = root.select_one(f'a[data-nlog-area$=".nav"][data-nlog-params*="{gdid}"]')
        body = root.select_one(f'a[data-nlog-area$=".body"][data-nlog-params*="{gdid}"]')

        is_head = title_a["data-nlog-area"].split(".")[-2] == "h"
        if is_head:
            group += 1  # 대표 기사가 나올 때마다 새 묶음이 시작된다

        # 언론사와 게시 시각은 네이버뉴스 링크가 들어 있는 Profile 블록에 있다
        profile = nav.find_parent(attrs={"data-sds-comp": "Profile"}) if nav else None
        press, list_time = "", ""
        if profile:
            press = text_of(profile.select_one(".sds-comps-profile-info-title-text"))
            for sub in profile.select(".sds-comps-profile-info-subtext"):
                if not sub.find("a"):  # 링크가 없는 쪽이 '7시간 전' 같은 시각
                    list_time = text_of(sub)
                    break

        naver_match = ARTICLE_URL.match(nav["href"]) if nav else None
        items.append({
            "rank": int(params.get("cr_rank_legacy") or len(items) + 1),
            "group": group,
            "type": "대표" if is_head else "관련",
            "press": press,
            "title": text_of(title_a),
            "snippet": text_of(body),
            "list_time": list_time,
            "published": "",
            "original_url": title_a.get("href", ""),
            "naver_url": naver_match.group(0) if naver_match else "",  # ?sid=... 는 버린다
            "content": "",
        })

    return sorted(items, key=lambda x: x["rank"])


# ---------- 2. 기사 본문 파싱 ----------

def parse_article(soup):
    """네이버 기사 페이지에서 제목, 작성 시각, 본문을 추출한다."""
    body = soup.select_one("#dic_area") or soup.select_one("#newsct_article")
    if body is None:
        raise ValueError("본문을 찾을 수 없습니다 (지원하지 않는 기사 형식일 수 있음)")

    for tag in body.select(NOISE_SELECTORS):
        if not tag.decomposed:  # 부모가 이미 제거된 경우는 건너뜀
            tag.decompose()
    for br in body.find_all("br"):
        br.replace_with("\n")  # 문단 구분은 <br>로 되어 있다

    lines = [re.sub(r"[ \t ]+", " ", line).strip() for line in body.get_text().splitlines()]
    date_el = soup.select_one(".media_end_head_info_datestamp_time")
    return {
        "title": text_of(soup.select_one("#title_area")),
        "published": (date_el.get("data-date-time") or text_of(date_el)) if date_el else "",
        "content": "\n".join(line for line in lines if line),
    }


# ---------- 3. 실행 ----------

def fetch_content(session, item, log=print):
    """네이버 기사 페이지에서 작성 시각과 본문을 가져와 item에 채운다. 성공하면 True."""
    if not item["naver_url"]:
        log("      (네이버뉴스 링크가 없어 본문은 건너뜁니다)")
        return False
    try:
        article = parse_article(get_soup(session, item["naver_url"]))
    except (requests.RequestException, ValueError) as e:
        log(f"      본문 수집 실패: {e}")
        return False
    item["published"] = article["published"]
    item["content"] = article["content"]
    if article["title"]:  # 목록의 제목은 '...'로 잘릴 수 있어 기사 페이지의 전체 제목을 쓴다
        item["title"] = article["title"]
    log(f"      본문 {len(item['content'])}자")
    return True


def crawl(search_url, limit, with_content, log=print, on_list=None, on_item=None, should_stop=None):
    """검색 목록을 파싱하고(필요하면 본문까지) 수집한다. GUI에서 재사용할 수 있게 콜백을 받는다.

    log(text)            진행 상황 메시지
    on_list(items)       목록 파싱 직후 호출 (본문은 아직 비어 있음)
    on_item(index, item) 항목 하나의 처리가 끝날 때마다 호출
    should_stop()        True를 돌려주면 다음 항목 처리 전에 중단
    """
    session = requests.Session()
    session.headers.update(HEADERS)

    items = parse_list(get_soup(session, search_url))[:limit]
    log(f"기사 {len(items)}개를 찾았습니다.\n")
    if on_list:
        on_list(items)

    for i, item in enumerate(items):
        if should_stop and should_stop():
            log("중지되었습니다.")
            break
        log(f"[{item['rank']:>2}] {item['type']} | {item['press']} | {item['list_time']} | {item['title']}")
        fetched = with_content and fetch_content(session, item, log)
        if on_item:
            on_item(i, item)
        if fetched and i < len(items) - 1:
            time.sleep(DELAY)
    return items


def save_json(items, path):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(items, f, ensure_ascii=False, indent=2)


def save_csv(items, path):
    # utf-8-sig: 엑셀에서 열어도 한글이 깨지지 않는다
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=FIELDS)
        writer.writeheader()
        writer.writerows(items)


XLSX_HEADERS = {
    "rank": "순위", "group": "묶음", "type": "구분", "press": "언론사", "title": "제목",
    "snippet": "목록 요약", "list_time": "목록 시각", "published": "작성 시각",
    "original_url": "원문 URL", "naver_url": "네이버뉴스 URL", "content": "본문",
}
XLSX_WIDTHS = {"rank": 6, "group": 6, "type": 6, "press": 14, "title": 50, "snippet": 50,
               "list_time": 12, "published": 20, "original_url": 40, "naver_url": 40, "content": 80}


def save_xlsx(items, path):
    """엑셀(.xlsx)로 저장한다. openpyxl이 필요하다: pip install openpyxl"""
    from openpyxl import Workbook
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "뉴스"
    ws.append([XLSX_HEADERS[f] for f in FIELDS])

    for item in items:
        row = []
        for field in FIELDS:
            value = item[field]
            if isinstance(value, str):  # 엑셀이 허용하지 않는 제어 문자 제거, 셀 한도(32,767자) 대비 자르기
                value = ILLEGAL_CHARACTERS_RE.sub("", value)[:32000]
            row.append(value)
        ws.append(row)

    header_fill = PatternFill("solid", fgColor="1F3864")
    for col, field in enumerate(FIELDS, 1):
        head = ws.cell(row=1, column=col)
        head.font = Font(bold=True, color="FFFFFF")
        head.fill = header_fill
        head.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col)].width = XLSX_WIDTHS[field]
        for row in range(2, len(items) + 2):
            cell = ws.cell(row=row, column=col)
            centered = field in ("rank", "group", "type", "list_time", "published")
            cell.alignment = Alignment(
                horizontal="center" if centered else "left", vertical="center", wrap_text=(field == "title"))
            if field in ("original_url", "naver_url") and cell.value:
                cell.hyperlink = cell.value
                cell.style = "Hyperlink"
    ws.freeze_panes = "A2"  # 머리글 고정
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def save(items, prefix):
    save_json(items, f"{prefix}.json")
    save_csv(items, f"{prefix}.csv")
    saved = [f"{prefix}.json", f"{prefix}.csv"]
    try:
        save_xlsx(items, f"{prefix}.xlsx")
        saved.append(f"{prefix}.xlsx")
    except ImportError:
        print("openpyxl이 없어 엑셀 파일은 만들지 못했습니다. (pip install openpyxl)")
    print(f"\n저장 완료: {', '.join(saved)}")


def main():
    parser = argparse.ArgumentParser(description="네이버 검색 결과의 뉴스 목록/본문 크롤러")
    parser.add_argument("-u", "--url", default=SEARCH_URL, help="네이버 검색 결과 URL (기본: '반도체')")
    parser.add_argument("-n", "--limit", type=int, default=10, help="수집할 기사 수 (기본: 10)")
    parser.add_argument("-o", "--output", default="news_articles", help="저장 파일 이름(확장자 제외)")
    parser.add_argument("--list-only", action="store_true", help="본문은 가져오지 않고 검색 목록만 수집")
    args = parser.parse_args()

    items = crawl(args.url, args.limit, with_content=not args.list_only)
    if not items:
        print("수집된 기사가 없습니다. 검색 결과의 태그 구조가 바뀌었는지 확인하세요.")
        return
    save(items, args.output)

    first = items[0]
    print("\n--- 첫 번째 기사 미리보기 ---")
    print(f"{first['title']}\n{first['press']} · {first['published'] or first['list_time']}\n{first['original_url']}\n")
    if first["snippet"]:
        print(f"[요약] {first['snippet']}\n")
    if first["content"]:
        print(first["content"][:300] + ("..." if len(first["content"]) > 300 else ""))


if __name__ == "__main__":
    main()
